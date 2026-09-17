from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt

md_path = Path(r"C:\Users\28427\Desktop\test\水果电商说明书.md")
out_path = Path(r"C:\Users\28427\Desktop\test\水果电商说明书.docx")

md_text = md_path.read_text(encoding="utf-8")
lines = md_text.splitlines()

doc = Document()

code_block = False

for raw_line in lines:
    line = raw_line.rstrip("\n")

    if line.strip().startswith("```"):
        code_block = not code_block
        continue

    if code_block:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        continue

    if line.startswith("# "):
        doc.add_heading(line[2:].strip(), level=1)
        continue
    if line.startswith("## "):
        doc.add_heading(line[3:].strip(), level=2)
        continue
    if line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=3)
        continue

    img_match = re.match(r"!\[(.*?)\]\((.*?)\)", line.strip())
    if img_match:
        alt = img_match.group(1).strip()
        path = img_match.group(2).strip()
        img_path = (md_path.parent / path).resolve()
        if img_path.exists():
            doc.add_paragraph(alt)
            doc.add_picture(str(img_path), width=Inches(5.5))
        else:
            doc.add_paragraph(f"图片：{alt} ({path})")
        continue

    if re.match(r"^\d+\.\s+", line):
        text = re.sub(r"^\d+\.\s+", "", line)
        doc.add_paragraph(text, style="List Number")
        continue

    if line.startswith("- "):
        doc.add_paragraph(line[2:].strip(), style="List Bullet")
        continue

    if not line.strip():
        doc.add_paragraph("")
        continue

    doc.add_paragraph(line)

doc.save(out_path)
print(f"Saved: {out_path}")
